"""出力：Excel（別表ごとのシート・計算根拠・指摘）と CSV（全欄）。

これは申告書の「ドラフト」です。国税庁の正式様式でも、電子申告（e-Tax）の仕様に適合したデータでもありません。
"""
from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

DRAFT_NOTE = '【ドラフト】このファイルは確認用です。国税庁の正式様式・電子申告データではありません。最終的な内容は税理士等が確認してください。'


def _safe(v: Any) -> Any:
    """CSV・Excel で数式として解釈されないようにする。"""
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + v
    return v


def _display(cell: dict | None) -> Any:
    if cell is None:
        return ''
    if cell['unit'] != 'yen':
        return cell.get('text') or ''
    v = cell['value']
    return '未確定' if v is None else v


def to_xlsx(company: str, period: str, result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = '概要'
    bold = Font(bold=True)
    warn = PatternFill('solid', fgColor='FFF4CC')
    ws.append([DRAFT_NOTE])
    ws['A1'].font = Font(bold=True, color='9C0006')
    ws.append(['法人名', _safe(company)])
    ws.append(['事業年度', period])
    s = result['summary']
    for label, key in (('所得金額又は欠損金額', 'income'), ('差引確定法人税額', 'corporate_final'),
                       ('差引確定地方法人税額', 'local_corporate_final'), ('差引確定防衛特別法人税額', 'defense_final'),
                       ('翌期へ繰り越す欠損金額', 'loss_carryforward'), ('計上すべき未払法人税等（参考）', 'recommended_tax_provision')):
        v = s.get(key)
        ws.append([label, '未確定' if v is None else v])
    ws.append(['重大エラー', result['counts']['critical']])
    ws.append(['要確認（未確認）', result['counts']['warning_open']])
    ws.append(['計算エンジンの版', result['engine_version']])
    ws.append(['入力のハッシュ', result['input_hash']])
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 24

    cells = result['cells']
    for sch in result['schedules']:
        sh = wb.create_sheet(sch['title'][:31])
        sh.append([f'{sch["title"]} {sch["form_name"]}'])
        sh['A1'].font = bold
        sh.append([DRAFT_NOTE])
        if sch['status'] != 'computed':
            sh.append([sch['status_note']])
        cols = [c['id'] for c in sch['columns']]
        extra = sorted({c for r in sch['rows'] for c in r['cells'] if c not in cols})
        header = ['行', '区分'] + [c['label'] for c in sch['columns']] + extra
        sh.append(header)
        for c in sh[sh.max_row]:
            c.font = bold
        for r in sch['rows']:
            vals = [r['line'], _safe(r['label'])] + [_display(cells.get(r['cells'].get(c))) for c in cols + extra]
            sh.append(vals)
            if any(v == '未確定' for v in vals):
                for c in sh[sh.max_row]:
                    c.fill = warn
        sh.column_dimensions['B'].width = 60
        for col in 'CDEFGHIJ':
            sh.column_dimensions[col].width = 16

    tr = wb.create_sheet('計算根拠')
    tr.append(['欄', '別表', '行', '列', '項目', '金額', '計算式', '参照した欄', '根拠', '出所', '上書きの理由', 'メモ'])
    for c in tr[1]:
        c.font = bold
    for k, c in cells.items():
        tr.append([k, c['schedule'], c['line'], c['col'], _safe(c['label']), _display(c), _safe(c['formula']),
                   ', '.join(c['refs']), _safe(' / '.join(c['basis'])), c['source'], _safe(c['override_reason']),
                   _safe(' / '.join(c['notes']))])
    tr.column_dimensions['E'].width = 40
    tr.column_dimensions['G'].width = 50

    fs = wb.create_sheet('指摘')
    fs.append(['段階', '分類', '内容', '詳しく', 'どう直すか', '確認済み'])
    for c in fs[1]:
        c.font = bold
    for f in result['findings']:
        fs.append([f['severity_label'], f['category'], _safe(f['title']), _safe(f['message']), _safe(f['action']),
                   '済' if f.get('acknowledged') else ''])
    for col, w in zip('ABCDEF', (10, 12, 50, 70, 60, 8)):
        fs.column_dimensions[col].width = w
    for row in fs.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')

    rs = wb.create_sheet('使用したルール')
    rs.append(['ルール', '版', '適用期間（開始）', '適用期間（終了）', '判定の基準日', '状態', '確認者', '確認日', '専門家レビュー', '根拠資料'])
    for c in rs[1]:
        c.font = bold
    for r in result['rules_used']:
        rs.append([r['title'], r['version'], r['valid_from'], r['valid_to'] or '', r['basis_label'], r['status_label'],
                   r['checked_by'], r['checked_at'], r['expert_review'], ' / '.join(f'{s["title"]} {s["url"]}' for s in r['sources'])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(result: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['欄', '別表', '行', '列', '項目', '金額', '計算式', '根拠', '出所'])
    for k, c in result['cells'].items():
        w.writerow([k, c['schedule'], c['line'], c['col'], _safe(c['label']), _display(c), _safe(c['formula']),
                    _safe(' / '.join(c['basis'])), c['source']])
    return ('﻿' + buf.getvalue()).encode('utf-8')
