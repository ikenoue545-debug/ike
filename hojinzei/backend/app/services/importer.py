"""CSV・Excel の取込：文字コードの判定、見出し行の検出、列の自動マッピング、行の読み取り。

列名や並びが違っても、見出しの語から列を推定する。推定した対応は利用者が確認・修正してから取り込む。
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from typing import Any, Optional

from .accounts_master import heading_section, normalize

TB_FIELDS: dict[str, dict[str, Any]] = {
    'account_name': {'label': '勘定科目', 'required': True,
                     'aliases': ['勘定科目', '科目', '科目名', '勘定科目名', '分類と勘定科目', '勘定', '項目']},
    'account_code': {'label': '科目コード', 'required': False, 'aliases': ['勘定科目コード', '科目コード', 'コード', '科目番号']},
    'opening': {'label': '期首残高', 'required': False,
                'aliases': ['期首残高', '期首', '前期繰越', '前期繰越残高', '開始残高', '繰越残高', '前月繰越', '期首残']},
    'debit': {'label': '借方金額', 'required': True, 'aliases': ['借方金額', '借方', '借方発生額', '当期借方', '借方合計', '借方発生']},
    'credit': {'label': '貸方金額', 'required': True, 'aliases': ['貸方金額', '貸方', '貸方発生額', '当期貸方', '貸方合計', '貸方発生']},
    'closing': {'label': '期末残高', 'required': True,
                'aliases': ['期末残高', '期末', '残高', '当期残高', '差引残高', '翌期繰越', '期末残', '当月残高']},
}
JOURNAL_FIELDS: dict[str, dict[str, Any]] = {
    'date': {'label': '取引日', 'required': True, 'aliases': ['取引日', '日付', '発生日', '伝票日付', '計上日']},
    'entry_no': {'label': '伝票番号', 'required': False, 'aliases': ['伝票番号', '仕訳番号', '番号', '伝票No']},
    'debit_account': {'label': '借方勘定科目', 'required': True, 'aliases': ['借方勘定科目', '借方科目']},
    'debit_amount': {'label': '借方金額', 'required': True, 'aliases': ['借方金額']},
    'debit_tax': {'label': '借方税区分', 'required': False, 'aliases': ['借方税区分']},
    'debit_partner': {'label': '借方取引先', 'required': False, 'aliases': ['借方取引先']},
    'debit_item': {'label': '借方品目', 'required': False, 'aliases': ['借方品目']},
    'debit_department': {'label': '借方部門', 'required': False, 'aliases': ['借方部門']},
    'credit_account': {'label': '貸方勘定科目', 'required': True, 'aliases': ['貸方勘定科目', '貸方科目']},
    'credit_amount': {'label': '貸方金額', 'required': True, 'aliases': ['貸方金額']},
    'credit_tax': {'label': '貸方税区分', 'required': False, 'aliases': ['貸方税区分']},
    'credit_partner': {'label': '貸方取引先', 'required': False, 'aliases': ['貸方取引先']},
    'credit_item': {'label': '貸方品目', 'required': False, 'aliases': ['貸方品目']},
    'credit_department': {'label': '貸方部門', 'required': False, 'aliases': ['貸方部門']},
    'memo': {'label': '摘要', 'required': False, 'aliases': ['摘要', '備考', 'メモ', '内容']},
}
FIELDS = {'trial_balance': TB_FIELDS, 'journal': JOURNAL_FIELDS}


class ImportErrorJa(Exception):
    pass


def decode(raw: bytes) -> tuple[str, str]:
    for enc in ('utf-8-sig', 'cp932'):
        try:
            return raw.decode(enc), ('UTF-8' if enc == 'utf-8-sig' else 'Shift_JIS（CP932）')
        except UnicodeDecodeError:
            continue
    raise ImportErrorJa('文字コードを判定できません（UTF-8 または Shift_JIS で保存してください）')


def read_table(filename: str, raw: bytes, sheet: Optional[str] = None) -> tuple[list[list[str]], dict]:
    """ファイルを行の配列にする。"""
    name = filename.lower()
    if name.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ImportErrorJa(f'Excel ファイルを開けません：{e}') from e
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        rows = [['' if c is None else _cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
        return _trim(rows), {'format': 'Excel', 'sheet': ws.title, 'sheets': wb.sheetnames, 'encoding': None}
    if name.endswith('.xls'):
        raise ImportErrorJa('古い Excel 形式（.xls）は読めません。.xlsx または CSV で保存し直してください')
    text, enc = decode(raw)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',\t;')
        delim = dialect.delimiter
    except csv.Error:
        delim = ','
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return _trim(rows), {'format': 'CSV', 'encoding': enc, 'delimiter': delim}


def _cell_text(c: Any) -> str:
    if isinstance(c, datetime):
        return c.date().isoformat()
    if isinstance(c, date):
        return c.isoformat()
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)


def _trim(rows: list[list[str]]) -> list[list[str]]:
    out = [[(c or '').strip() for c in r] for r in rows]
    while out and not any(out[-1]):
        out.pop()
    return out


def _norm_header(s: str) -> str:
    s = unicodedata.normalize('NFKC', s or '')
    return re.sub(r'[\s　()（）\[\]［］:：]', '', s).replace('(円)', '').replace('円', '')


def detect(kind: str, rows: list[list[str]]) -> dict:
    """見出し行と列の対応を推定する。"""
    fields = FIELDS[kind]
    best = (-1, 0, {})
    for i, row in enumerate(rows[:30]):
        mapping = _map_row(fields, row)
        score = sum(2 if fields[f]['required'] else 1 for f in mapping)
        if score > best[0]:
            best = (score, i, mapping)
    score, header_row, mapping = best
    if score <= 0:
        raise ImportErrorJa('見出しの行が見つかりません（「勘定科目」「借方」などの列名が必要です）')
    header = rows[header_row] if rows else []
    # freee の試算表は「分類と勘定科目」の列名が空のことがある：科目の列が見つからなければ、文字の多い列を候補にする
    if kind == 'trial_balance' and 'account_name' not in mapping:
        cand = _guess_text_column(rows[header_row + 1:header_row + 40], exclude=set(mapping.values()))
        if cand is not None:
            mapping['account_name'] = cand
    missing = [fields[f]['label'] for f, d in fields.items() if d['required'] and f not in mapping]
    return {
        'header_row': header_row,
        'header': header,
        'mapping': mapping,
        'missing_required': missing,
        'fields': [{'key': k, 'label': d['label'], 'required': d['required']} for k, d in fields.items()],
        'sample': rows[header_row + 1:header_row + 11],
        'total_rows': max(len(rows) - header_row - 1, 0),
    }


def _map_row(fields: dict, row: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used: set[int] = set()
    normed = [_norm_header(c) for c in row]
    # 完全一致を優先し、次に前方一致
    for exact in (True, False):
        for f, d in fields.items():
            if f in mapping:
                continue
            for alias in d['aliases']:
                a = _norm_header(alias)
                for j, h in enumerate(normed):
                    if j in used or not h:
                        continue
                    if (h == a) if exact else (h.startswith(a) and len(h) <= len(a) + 4):
                        mapping[f] = j
                        used.add(j)
                        break
                if f in mapping:
                    break
    return mapping


def _guess_text_column(rows: list[list[str]], exclude: set[int]) -> Optional[int]:
    counts: dict[int, int] = {}
    for r in rows:
        for j, c in enumerate(r):
            if j in exclude or not c:
                continue
            if parse_amount(c) is None:
                counts[j] = counts.get(j, 0) + 1
    return max(counts, key=counts.get) if counts else None


_AMOUNT_RE = re.compile(r'^[\-−▲△]?\(?[\d,，]+(\.\d+)?\)?$')


def parse_amount(s: str) -> Optional[int]:
    """'1,234' '△1,234' '(1,234)' '-1234' '¥1,234円' を整数円にする。空欄は None。"""
    if s is None:
        return None
    t = unicodedata.normalize('NFKC', str(s)).strip().replace('¥', '').replace('円', '').replace(' ', '')
    if t in ('', '-', '—', '―'):
        return None
    neg = False
    if t[0] in '-−▲△':
        neg, t = True, t[1:]
    if t.startswith('(') and t.endswith(')'):
        neg, t = True, t[1:-1]
    t = t.replace(',', '').replace('，', '')
    if not re.fullmatch(r'\d+(\.\d+)?', t):
        return None
    if '.' in t:
        whole, frac = t.split('.')
        if int(frac or 0) != 0:
            raise ImportErrorJa(f'円未満の金額は扱えません：{s}')
        t = whole
    v = int(t)
    return -v if neg else v


def parse_date(s: str) -> Optional[date]:
    t = unicodedata.normalize('NFKC', s or '').strip()
    if not t:
        return None
    for fmt in ('%Y/%m/%d', '%Y-%m-%d', '%Y.%m.%d', '%Y年%m月%d日'):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    m = re.match(r'^(\d{4})(\d{2})(\d{2})$', t)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    return None


_SUBTOTAL_RE = re.compile(r'(合計|小計|計)$')


def parse_trial_balance(rows: list[list[str]], header_row: int, mapping: dict[str, int]) -> dict:
    """試算表の行を読む。見出し（流動資産など）・小計の行は科目にしない。"""
    def cell(r, f):
        j = mapping.get(f)
        return r[j] if j is not None and j < len(r) else ''

    accounts, skipped, errors = [], [], []
    heading: Optional[str] = None
    seen: set[str] = set()
    for i, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        name = cell(r, 'account_name').strip()
        if not name:
            continue
        try:
            opening, debit, credit, closing = (parse_amount(cell(r, f)) for f in ('opening', 'debit', 'credit', 'closing'))
        except ImportErrorJa as e:
            errors.append({'row': i, 'message': str(e)})
            continue
        nn = normalize(name)
        if _SUBTOTAL_RE.search(nn):
            skipped.append({'row': i, 'label': name, 'reason': '小計・合計の行'})
            continue
        if heading_section(name) is not None and debit is None and credit is None and closing is None:
            heading = name
            skipped.append({'row': i, 'label': name, 'reason': '見出しの行'})
            continue
        if nn in ('差引損益計算', '当期純損益金額', '当期純利益', '当期純損失', '当期純損益', '税引前当期純利益', '経常利益',
                  '営業利益', '売上総利益', '売上総損益金額', '営業損益金額', '経常損益金額', '税引前当期純損益金額'):
            skipped.append({'row': i, 'label': name, 'reason': '利益の行（計算で求める）'})
            continue
        if debit is None and credit is None and closing is None and opening is None:
            heading = name
            skipped.append({'row': i, 'label': name, 'reason': '金額のない行（見出しとして扱う）'})
            continue
        if nn in seen:
            errors.append({'row': i, 'message': f'科目「{name}」が2回あります（後の行を無視しました）'})
            continue
        seen.add(nn)
        accounts.append({'row': i, 'name': name, 'code': cell(r, 'account_code') or None, 'heading': heading,
                         'opening': opening, 'debit': debit, 'credit': credit, 'closing': closing})
    return {'accounts': accounts, 'skipped': skipped, 'errors': errors}


def parse_journal(rows: list[list[str]], header_row: int, mapping: dict[str, int]) -> dict:
    def cell(r, f):
        j = mapping.get(f)
        return r[j] if j is not None and j < len(r) else ''

    lines, errors = [], []
    last_date, last_no = None, None
    for i, r in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not any(r):
            continue
        d = parse_date(cell(r, 'date')) or last_date
        no = cell(r, 'entry_no') or last_no
        if cell(r, 'date'):
            last_date, last_no = d, no
        try:
            for side in ('debit', 'credit'):
                acc = cell(r, f'{side}_account').strip()
                amt = parse_amount(cell(r, f'{side}_amount'))
                if not acc or amt is None:
                    continue
                other = 'credit' if side == 'debit' else 'debit'
                lines.append({'row': i, 'entry_date': d, 'entry_no': no, 'side': side, 'account_name': acc,
                              'counter_account': cell(r, f'{other}_account').strip() or None, 'amount': amt,
                              'tax_category': cell(r, f'{side}_tax') or None, 'partner': cell(r, f'{side}_partner') or None,
                              'item': cell(r, f'{side}_item') or None, 'department': cell(r, f'{side}_department') or None,
                              'memo': cell(r, 'memo') or None})
        except ImportErrorJa as e:
            errors.append({'row': i, 'message': str(e)})
        if d is None:
            errors.append({'row': i, 'message': '日付を読めません'})
    return {'lines': lines, 'errors': errors}
